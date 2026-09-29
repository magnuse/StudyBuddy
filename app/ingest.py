"""Turns files that parents send into plain text."""

from __future__ import annotations

import io
import logging
from pathlib import Path

log = logging.getLogger(__name__)

IMAGE_TYPES = {".jpg", ".jpeg", ".png", ".webp", ".heic", ".heif", ".gif", ".bmp", ".tif", ".tiff"}
OCR_LANGUAGES = "swe+spa+eng"


class UnsupportedFile(Exception):
    pass


def extract_text(data: bytes, filename: str) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".pdf":
        return _pdf(data)
    if suffix == ".docx":
        return _docx(data)
    if suffix in {".txt", ".md", ".csv"}:
        return data.decode("utf-8", errors="replace")
    if suffix in IMAGE_TYPES:
        return _ocr(data)
    if suffix in {".doc", ".ppt", ".odt"}:
        raise UnsupportedFile("Gamla Office-format stöds inte än. Spara som PDF eller skicka en skärmbild.")
    raise UnsupportedFile(f"Filtypen {suffix or 'okänd'} stöds inte. Skicka PDF, Word, bild eller text.")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for number, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[Sida {number}]\n{text}")
    if not pages:
        raise UnsupportedFile("PDF:en innehåller ingen text (den är troligen inskannad). Skicka sidorna som bilder i stället.")
    return "\n\n".join(pages)


def _docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" - ".join(cells))
    return "\n".join(lines)


def _ocr(data: bytes) -> str:
    from PIL import Image

    try:
        import pillow_heif

        pillow_heif.register_heif_opener()
    except ImportError:
        pass
    try:
        import pytesseract
    except ImportError as exc:
        raise UnsupportedFile("Bildläsning (OCR) är inte installerad på servern.") from exc

    image = Image.open(io.BytesIO(data))
    image = image.convert("L")  # greyscale reads better
    text = pytesseract.image_to_string(image, lang=OCR_LANGUAGES)
    if not text.strip():
        raise UnsupportedFile("Kunde inte läsa någon text i bilden. Prova en skarpare bild i bra ljus.")
    return text
